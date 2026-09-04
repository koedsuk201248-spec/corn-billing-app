import streamlit as st
import openpyxl
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import io

# ตั้งค่าหน้าเว็บ Streamlit
st.set_page_config(page_title="ระบบสร้างใบวางบิลข้าวโพดแห้ง", page_icon="🌽", layout="centered")

def add_highlight(run):
    rPr = run._r.get_or_add_rPr()
    highlight = OxmlElement('w:highlight')
    highlight.set(qn('w:val'), 'yellow')
    rPr.append(highlight)

st.title("🌽 ระบบสร้างใบวางบิลข้าวโพดแห้ง (เจ้นัชชา)")
st.write("เลือกอัปโหลดไฟล์ Excel แล้วติ๊กเลือกข้อที่ต้องการนำมาสร้างเป็นไฟล์ Word ได้ทันที")

# 1. ปุ่มอัปโหลดไฟล์ Excel
uploaded_file = st.file_uploader("อัปโหลดไฟล์ Excel (ขายข้าวโพดแห้ง)", type=["xlsx"])

if uploaded_file is not None:
    wb = openpyxl.load_workbook(uploaded_file, data_only=True)
    
    # 2. รายชื่อชีทจริงในไฟล์ Excel
    raw_sheet_names = wb.sheetnames
    
    # แสดงชื่อชีทแบบตัดเว้นวรรคเพื่อความสวยงามใน Dropdown
    sheet_options = {s.strip(): s for s in raw_sheet_names}
    
    selected_display = st.selectbox("📌 เลือกอำเภอ / ชีทที่ต้องการ:", list(sheet_options.keys()))
    
    if selected_display:
        # ดึงชื่อชีทจริงใน Excel มาใช้งาน
        actual_sheet_name = sheet_options[selected_display]
        ws = wb[actual_sheet_name]
        
        # อ่านรายการข้อมูลในชีทที่เลือก
        items_data = []
        for row in range(4, ws.max_row + 1):
            plate = ws.cell(row=row, column=4).value
            destination = ws.cell(row=row, column=8).value
            weight_end = ws.cell(row=row, column=11).value
            weight_diff = ws.cell(row=row, column=12).value
            price = ws.cell(row=row, column=13).value
            
            if plate and str(plate).strip() != "ทะเบียน":
                items_data.append({
                    "plate": str(plate).strip(),
                    "destination": str(destination).strip() if destination else "-",
                    "weight_end": weight_end,
                    "weight_diff": weight_diff,
                    "price": price
                })
        
        st.write("---")
        st.write(f"### 📋 รายการข้อมูลอำเภอ **{selected_display}** (พบทั้งหมด {len(items_data)} รายการ)")
        
        if len(items_data) == 0:
            st.warning("⚠️ ไม่พบข้อมูลรายการในอำเภอนี้")
        else:
            # 3. เลือกข้อที่ต้องการนำมาทำ Word
            st.write("ติ๊กเลือกข้อที่ต้องการทำ Word:")
            
            col_a, col_b = st.columns(2)
            select_all = col_a.button("✅ เลือกทั้งหมด")
            clear_all = col_b.button("❌ ไม่เลือกเลย")
            
            selected_indices = []
            
            for idx, item in enumerate(items_data, 1):
                label = f"ข้อ {idx}: ทะเบียน {item['plate']} | ปลายทาง: {item['destination']} | นน.: {item['weight_end'] or '-'} | ราคา: {item['price'] or '-'}"
                
                default_val = True
                if clear_all:
                    default_val = False
                elif select_all:
                    default_val = True
                    
                is_selected = st.checkbox(label, value=default_val, key=f"chk_{idx}")
                if is_selected:
                    selected_indices.append(idx - 1)
            
            st.write("---")
            
            # 4. ปุ่มสร้างและดาวน์โหลดไฟล์ Word
            if st.button("🚀 สร้างไฟล์ Word (.docx)", type="primary"):
                if len(selected_indices) == 0:
                    st.error("กรุณาเลือกอย่างน้อย 1 ข้อก่อนสั่งสร้างไฟล์ครับ")
                else:
                    doc = Document()
                    
                    # ตั้งค่าระยะขอบ A4
                    for section in doc.sections:
                        section.page_width = Inches(8.27)
                        section.page_height = Inches(11.69)
                        section.top_margin = Inches(1.0)
                        section.bottom_margin = Inches(1.0)
                        section.left_margin = Inches(1.2)
                        section.right_margin = Inches(1.2)
                        
                    p_title = doc.add_paragraph()
                    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_title.paragraph_format.space_after = Pt(24)
                    r_title = p_title.add_run(f"วางบิลข้าวโพดแห้ง เจ้นัชชา ({selected_display})")
                    r_title.font.name = "TH Sarabun PSK"
                    r_title.font.size = Pt(22)
                    r_title.font.bold = True
                    
                    for count, s_idx in enumerate(selected_indices, 1):
                        item = items_data[s_idx]
                        
                        if count > 1:
                            doc.add_page_break()
                            
                        header_text = f"ทะเบียน {item['plate']} ({selected_display}-{item['destination']})"
                        
                        weight_val = item['weight_end']
                        weight_str = f"นน.ปลายทาง {weight_val:,.0f}" if isinstance(weight_val, (int, float)) else f"นน.ปลายทาง {weight_val or '-'}"
                        
                        diff_val = item['weight_diff']
                        if isinstance(diff_val, (int, float)):
                            sign = "+" if diff_val > 0 else ""
                            diff_str = f"ส่วนต่างน้ำหนัก {sign}{diff_val:,.0f} กิโลกรัม"
                        else:
                            diff_str = f"ส่วนต่างน้ำหนัก {diff_val or '-'}"
                            
                        price_str = f"ราคา {item['price']}" if item['price'] else "ราคา -"
                        
                        lines = [header_text, weight_str, diff_str, price_str]
                        for text in lines:
                            p = doc.add_paragraph()
                            p.paragraph_format.space_after = Pt(12)
                            r = p.add_run(text)
                            r.font.name = "TH Sarabun PSK"
                            r.font.size = Pt(16)
                            
                        p_status = doc.add_paragraph()
                        p_status.paragraph_format.space_after = Pt(12)
                        r_status = p_status.add_run("*ลงเรียบร้อย*")
                        r_status.font.name = "TH Sarabun PSK"
                        r_status.font.size = Pt(16)
                        r_status.font.bold = True
                        add_highlight(r_status)
                    
                    # เซฟไฟล์เข้าสตรีมบัฟเฟอร์เตรียมโหลด
                    bio = io.BytesIO()
                    doc.save(bio)
                    bio.seek(0)
                    
                    st.success(f"สร้างไฟล์สำเร็จ! สรุปเลือกมาทำทั้งหมด {len(selected_indices)} ข้อ")
                    st.download_button(
                        label="📥 กดดาวน์โหลดไฟล์ Word (.docx)",
                        data=bio,
                        file_name=f"วางบิลข้าวโพดแห้ง_{selected_display}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"   
                    )